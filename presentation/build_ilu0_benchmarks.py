"""Build an editable slide deck from the saved ILU0 row-dependency results."""

from __future__ import annotations

import csv
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
from scipy.io import mmread


ROOT = Path(__file__).resolve().parents[1]
RESULTS = ROOT / "ilu0-parallel-row-dependencies.csv"
MATRIX = ROOT / "benchmark/matrices/suitesparse/bcsstk36/bcsstk36.mtx"
OUTPUT = Path(__file__).resolve().parent / "ilu0_benchmarks.pptx"

NAVY = RGBColor(12, 27, 46)
PANEL = RGBColor(22, 43, 67)
PANEL_LIGHT = RGBColor(29, 55, 80)
WHITE = RGBColor(247, 250, 252)
MUTED = RGBColor(170, 190, 207)
TEAL = RGBColor(78, 215, 198)
ORANGE = RGBColor(255, 183, 95)
RED = RGBColor(248, 119, 115)
BLUE = RGBColor(105, 178, 250)

WIDTH = 13.333
HEIGHT = 7.5


def rectangle(slide, x, y, w, h, fill, radius=False, line=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background() if line is None else None
    if line is not None:
        shape.line.color.rgb = line
    if radius:
        try:
            shape.adjustments[0] = 0.12
        except (AttributeError, IndexError):
            pass
    return shape


def label(slide, x, y, w, h, text, size=18, color=WHITE, bold=False,
          align=PP_ALIGN.LEFT, font="Aptos", valign=MSO_ANCHOR.MIDDLE):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.02)
    frame.margin_right = Inches(0.02)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def base_slide(prs, number, title, kicker=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rectangle(slide, 0, 0, WIDTH, HEIGHT, NAVY)
    rectangle(slide, 0.5, 0.48, 0.08, 0.49, TEAL)
    if kicker:
        label(slide, 0.77, 0.35, 11.7, 0.25, kicker.upper(), 10, TEAL, True)
    label(slide, 0.77, 0.73 if kicker else 0.51, 11.85, 0.72, title, 29, WHITE, True)
    rectangle(slide, 0.65, 7.12, 12.04, 0.015, PANEL_LIGHT)
    label(slide, 0.7, 7.16, 10.9, 0.2,
          "Source: local ILU0 benchmark code and saved SuiteSparse bcsstk36 results",
          8, MUTED)
    label(slide, 12.05, 7.13, 0.5, 0.27, f"{number:02}", 9, MUTED,
          align=PP_ALIGN.RIGHT)
    return slide


def card(slide, x, y, w, h, title, body, accent=TEAL, body_size=17):
    rectangle(slide, x, y, w, h, PANEL, radius=True)
    rectangle(slide, x + 0.2, y + 0.25, 0.06, 0.48, accent)
    label(slide, x + 0.37, y + 0.2, w - 0.6, 0.58, title, 20, WHITE, True)
    label(slide, x + 0.31, y + 0.91, w - 0.61, h - 1.08, body,
          body_size, MUTED, valign=MSO_ANCHOR.TOP)


def circle(slide, x, y, text, fill=BLUE, size=0.62):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y),
                                   Inches(size), Inches(size))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    label(slide, x, y + 0.01, size, size - 0.02, text, 20, NAVY, True,
          align=PP_ALIGN.CENTER)



def dependency_line(slide, x1, y1, x2, y2):
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1), Inches(y1), Inches(x2), Inches(y2),
    )
    connector.line.color.rgb = MUTED
    connector.line.width = Pt(1.8)


def source_data():
    with RESULTS.open(newline="", encoding="utf-8") as file:
        rows = list(csv.DictReader(file))
    selected = next((row for row in rows if "bcsstk36" in row["matrices"]), None)
    if selected is None:
        raise ValueError("Saved results do not contain bcsstk36")
    matrix = mmread(MATRIX)
    n = matrix.shape[0]
    if matrix.shape[0] != matrix.shape[1]:
        raise ValueError("Benchmark matrix must be square")
    names = {
        "Wavefront E2E": "Wavefront E2E (Rows/s)",
        "Wavefront Prepared": "Wavefront Prepared (Rows/s)",
        "Sync-free E2E": "Sync-free E2E (Rows/s)",
        "Sync-free Prepared": "Sync-free Prepared (Rows/s)",
        "CPU": "CPU Row Dependencies (Rows/s)",
    }
    milliseconds = {name: n * 1000 / float(selected[column])
                    for name, column in names.items()}
    return n, matrix.nnz, milliseconds


def main():
    n, nnz, times = source_data()
    prs = Presentation()
    prs.slide_width = Inches(WIDTH)
    prs.slide_height = Inches(HEIGHT)
    prs.core_properties.title = "ILU0 dependency scheduling benchmarks"
    prs.core_properties.subject = "Wavefront versus sync-free on SuiteSparse bcsstk36"

    # 1. Title
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rectangle(slide, 0, 0, WIDTH, HEIGHT, NAVY)
    rectangle(slide, 0.75, 1.08, 0.1, 4.62, TEAL)
    label(slide, 1.1, 1.13, 10.7, 0.42, "GPU SPARSE COMPUTATION", 17, TEAL, True)
    label(slide, 1.1, 1.81, 11.0, 1.56, "ILU0 dependency\nscheduling benchmarks",
          42, WHITE, True)
    label(slide, 1.1, 3.8, 10.7, 0.63,
          "Wavefront and sync-free execution on SuiteSparse bcsstk36", 22, MUTED)
    rectangle(slide, 1.1, 5.16, 3.65, 0.47, PANEL, radius=True)
    label(slide, 1.26, 5.2, 3.37, 0.37,
          "Row-dependency kernel study", 14, WHITE, True)
    label(slide, 1.1, 6.78, 10.9, 0.3,
          "Based on the repository's saved benchmark results", 11, MUTED)

    # 2. Scope
    slide = base_slide(prs, 2, "What was measured?", "Scope")
    card(slide, 0.75, 1.75, 5.76, 3.98, "Each row waits for earlier rows",
         "For every lower-triangular entry j < i, row i reads the result of row j. "
         "That gives a directed edge j → i in the dependency graph.", BLUE, 20)
    rectangle(slide, 6.78, 1.75, 5.76, 3.98, PANEL, radius=True)
    label(slide, 7.13, 2.13, 4.94, 0.4, "Benchmark loop", 20, WHITE, True)
    label(slide, 7.13, 2.75, 5.0, 1.25,
          "row_done[i] = 1 + Σ row_done[j]\n                   for j < i in row i",
          22, TEAL, font="Aptos Mono")
    label(slide, 7.13, 4.42, 4.97, 0.88,
          "This is a dependency surrogate. It does not compute the numerical ILU0 factors.",
          17, ORANGE)
    label(slide, 0.85, 6.27, 11.7, 0.45,
          "The study isolates scheduling and DAG preparation costs.", 20, WHITE, True)

    # 3. One DAG, two schedulers
    slide = base_slide(prs, 3, "One dependency graph, two schedulers", "Method")
    label(slide, 0.87, 1.6, 5.2, 0.43, "Wavefront: synchronized batches", 21, BLUE, True)
    rectangle(slide, 0.75, 2.12, 5.83, 3.39, PANEL, radius=True)
    for y, names in [(2.52, "A      B"), (3.45, "C      D"), (4.38, "E")]:
        rectangle(slide, 1.16, y, 4.96, 0.58, PANEL_LIGHT, radius=True)
        label(slide, 1.39, y + 0.03, 4.35, 0.48, names, 18, WHITE, True,
              align=PP_ALIGN.CENTER)
    label(slide, 1.17, 5.73, 5.2, 0.77,
          "Launch a wave → wait for all rows → launch the next wave", 17, MUTED)
    label(slide, 7.05, 1.6, 5.2, 0.43, "Sync-free: ready rows proceed", 21, TEAL, True)
    rectangle(slide, 6.75, 2.12, 5.83, 3.39, PANEL, radius=True)
    for edge in [(7.83, 3.27, 9.29, 3.92),
                 (7.83, 3.30, 10.45, 3.92),
                 (9.07, 3.30, 10.45, 3.92),
                 (9.91, 4.20, 11.08, 4.76),
                 (11.05, 4.23, 11.36, 4.46)]:
        dependency_line(slide, *edge)
    for x, y, text, color in [(7.25, 2.77, "A", BLUE), (8.49, 2.77, "B", BLUE),
                               (9.29, 3.61, "C", TEAL), (10.45, 3.61, "D", TEAL),
                               (11.06, 4.46, "E", ORANGE)]:
        circle(slide, x, y, text, color)
    label(slide, 7.16, 5.73, 5.2, 0.77,
          "One launch; atomic counters release each row when ready", 17, MUTED)
    label(slide, 1.05, 6.64, 11.2, 0.35,
          "‘Sync-free’ removes global wave barriers; rows still coordinate through atomic counters.",
          15, WHITE, align=PP_ALIGN.CENTER)

    # 4. Timing definitions
    slide = base_slide(prs, 4, "How to read the five methods", "Benchmark design")
    card(slide, 0.7, 1.65, 3.87, 3.94, "GPU end to end",
         "Build the DAG and schedule, prepare device arrays, execute on GPU, and copy the result back.",
         BLUE, 19)
    card(slide, 4.73, 1.65, 3.87, 3.94, "GPU prepared",
         "DAG and device arguments are prepared before timing. Sync-free resets its counters on every measured run.",
         TEAL, 19)
    card(slide, 8.76, 1.65, 3.87, 3.94, "CPU",
         "Execute the same row-dependency loop sequentially. No GPU setup or transfer is needed.",
         ORANGE, 19)
    label(slide, 0.93, 6.08, 11.55, 0.43,
          "Prepared = steady-state execution; end to end = one run on a new sparsity pattern.",
          20, WHITE, True, align=PP_ALIGN.CENTER)

    # 5. Dataset and measurements
    slide = base_slide(prs, 5, "Measured latency on bcsstk36", "Results · lower is better")
    label(slide, 0.8, 1.58, 11.6, 0.4,
          f"SuiteSparse matrix  ·  {n:,} × {n:,}  ·  {nnz:,} nonzeros",
          17, MUTED)
    methods = [
        ("Wavefront E2E", BLUE, "GPU"),
        ("Wavefront Prepared", BLUE, "GPU"),
        ("Sync-free E2E", TEAL, "GPU"),
        ("Sync-free Prepared", TEAL, "GPU"),
        ("CPU", ORANGE, "CPU"),
    ]
    max_ms = max(times.values())
    for index, (name, color, device) in enumerate(methods):
        y = 2.12 + index * 0.78
        label(slide, 0.82, y, 3.1, 0.36, name, 17, WHITE, True)
        rectangle(slide, 3.98, y + 0.04, 6.5, 0.36, PANEL_LIGHT, radius=True)
        rectangle(slide, 3.98, y + 0.04, 6.5 * times[name] / max_ms, 0.36,
                  color, radius=True)
        label(slide, 10.7, y - 0.01, 1.02, 0.46, f"{times[name]:.1f} ms",
              17, WHITE, True, align=PP_ALIGN.RIGHT)
        label(slide, 11.85, y - 0.01, 0.66, 0.46, device, 12, color, True,
              align=PP_ALIGN.RIGHT)
    label(slide, 0.82, 6.42, 11.5, 0.36,
          "Times reconstructed from saved rows/s; the benchmark reports the median.",
          13, MUTED)

    # 6. One-shot comparison
    slide = base_slide(prs, 6, "For one new matrix pattern, CPU wins", "One-shot decision")
    rectangle(slide, 0.88, 1.87, 11.58, 3.78, PANEL, radius=True)
    label(slide, 1.24, 2.3, 2.3, 0.4, "CPU", 22, ORANGE, True)
    rectangle(slide, 3.58, 2.35, 5.5 * times["CPU"] / times["Sync-free E2E"],
              0.52, ORANGE, radius=True)
    label(slide, 9.47, 2.31, 2.3, 0.54, f"{times['CPU']:.1f} ms", 23,
          WHITE, True)
    label(slide, 1.24, 3.56, 2.3, 0.77, "Sync-free\nGPU E2E", 21, TEAL, True)
    rectangle(slide, 3.58, 3.72, 5.5, 0.52, TEAL, radius=True)
    label(slide, 9.47, 3.69, 2.3, 0.54,
          f"{times['Sync-free E2E']:.1f} ms", 23, WHITE, True)
    label(slide, 1.24, 4.93, 10.9, 0.41,
          f"The GPU end-to-end path takes {times['Sync-free E2E']/times['CPU']:.1f}× as long.",
          20, WHITE, True)
    label(slide, 1.06, 6.15, 11.22, 0.58,
          "DAG construction, allocation, transfer, and result copying are paid on every one-shot run.",
          18, MUTED, align=PP_ALIGN.CENTER)

    # 7. Reuse illustration
    slide = base_slide(prs, 7, "Preparation can pay off when reused", "Repeated pattern")
    setup = times["Sync-free E2E"] - times["Sync-free Prepared"]
    label(slide, 0.86, 1.59, 11.6, 0.53,
          f"Approximate initial cost: {setup:.1f} ms  ·  prepared run: {times['Sync-free Prepared']:.1f} ms",
          19, MUTED)
    rectangle(slide, 0.82, 2.28, 11.72, 3.47, PANEL, radius=True)
    label(slide, 1.18, 2.61, 4.2, 0.4, "Repeated runs", 20, WHITE, True)
    for i, k in enumerate([1, 2, 3, 4, 10]):
        x = 1.2 + 2.22 * i
        cpu = k * times["CPU"]
        gpu = setup + k * times["Sync-free Prepared"]
        rectangle(slide, x, 3.16, 1.98, 2.08,
                  PANEL_LIGHT if k != 4 else RGBColor(28, 83, 87), radius=True)
        label(slide, x + 0.07, 3.35, 1.84, 0.37, f"{k} run{'s' if k > 1 else ''}",
              18, WHITE, True, align=PP_ALIGN.CENTER)
        label(slide, x + 0.08, 3.91, 1.79, 0.37,
              f"CPU  {cpu:.0f} ms", 14, ORANGE, True)
        label(slide, x + 0.08, 4.36, 1.79, 0.37,
              f"GPU  {gpu:.0f} ms", 14, TEAL, True)
    label(slide, 1.0, 6.02, 11.34, 0.68,
          "Illustrative break-even: ~4 runs with the same sparsity pattern. "
          "This estimate subtracts separate benchmark measurements; measure a true repeated-run workflow next.",
          15, MUTED, align=PP_ALIGN.CENTER)

    # 8. Conclusion
    slide = base_slide(prs, 8, "What the benchmark supports", "Takeaways")
    card(slide, 0.73, 1.77, 3.88, 3.88, "1 · Scheduler matters",
         f"Prepared sync-free is {times['Wavefront Prepared']/times['Sync-free Prepared']:.1f}× faster "
         "than prepared wavefront on bcsstk36.", TEAL, 20)
    card(slide, 4.73, 1.77, 3.88, 3.88, "2 · Setup matters",
         "The CPU wins a one-shot run. Reusing the same sparsity pattern can amortize preparation.",
         ORANGE, 20)
    card(slide, 8.73, 1.77, 3.88, 3.88, "3 · Next experiment",
         "Benchmark numerical ILU0 itself, measure the preparation stages separately, and test more real matrices.",
         BLUE, 20)
    label(slide, 0.92, 6.11, 11.45, 0.52,
          "Claim this as a row-dependency scheduling result, not a full ILU0 speedup.",
          20, WHITE, True, align=PP_ALIGN.CENTER)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    print(f"Saved {OUTPUT} ({len(prs.slides)} slides)")
    for name, milliseconds in times.items():
        print(f"{name}: {milliseconds:.3f} ms")


if __name__ == "__main__":
    main()
